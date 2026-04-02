from io import BytesIO

from openpyxl import load_workbook
from rest_framework.views import APIView 
from rest_framework.response import Response
from django.http import HttpResponse
from rest_framework.permissions import AllowAny, IsAuthenticated
from app.Models.events import Events
from app.Models.event_posts import EventPosts
from app.Models.event_attendees import EventAttendees
# from app.Models.event_badges_earned import EventBadgesEarned
from app.Models.venues_participating_in_event import VenuesParticipatingEvents
from app.Models.events import Events
from app.Models.challenge_achiever import Challenge_Achiever
from venue.Serializers.venue_info_serializer import VenueInfoSerializer
from venue.models.badges import BadgesLevel
from app.Models.event_posts import EventPosts
from venue.models.venue_info import Venue_Info
from app.Serializers.badge_level_serializer import BadgesLevelSerializer
from organisers.serializers.events_attendees_serializer import EventAppAttendeesSerializer
from organisers.serializers.events_posts_serializer import EventAppPostsSerializer
from organisers.serializers.events_venue_participating import EventAppVenuesParticipatingSerializer
from organisers.serializers.events_serializer import EventStaffSerializer
from staff.Permissions.adminOrganiserOnlyPermission import Request_By_Admin_And_Organiser_Only 
from rest_framework import filters
from django.utils import timezone
from datetime import timedelta
from django.db.models.functions import TruncMonth
from django.db.models import Count, Q
from app.Serializers.challenge_achiever_serializer import ChallengeAchieverSerializer
import pandas as pd


class EventStatsReportExportView(APIView):
    permission_classes = [AllowAny]

    def get(self, request):
        Events_qs = Events.objects.all().order_by('-id')
        now = timezone.now()

        data = []
        event_row_ranges = []
        row_start = 0  # 0-based index

        for event in Events_qs:
            participants = VenuesParticipatingEvents.objects.filter(event=event)
            attendees = EventAttendees.objects.filter(
                event=event, joined_at__lte=now
            ).count() or 0

            total_badges_earned = 0
            badge_required_to_earn = 0

            event_rows = []

            # ✅ If no participants
            if not participants.exists():
                event_rows.append({
                    "Event Title": event.title,
                    "Total Attendees": attendees,
                    "Participating Venues": 0,
                    "Badge Completion Rates": 100,
                    "Venue Name": "No Participating Venues",
                    "Total Scans": 0,
                    "Unique Visitors": 0,
                })

            else:
                # 🔹 Step 1: calculate totals
                for venue in participants:
                    scans_per_venue = Challenge_Achiever.objects.filter(
                        challenge__venue=venue.venues,
                        scanned_at__gte=event.event_start_date,
                        scanned_at__lte=event.event_close_date
                    )
                    total_badges_earned += scans_per_venue.count()
                    badge_required_to_earn += venue.scans_to_achieve_next_tier

                # 🔹 Step 2: calculate badge completion rate
                if badge_required_to_earn > 0:
                    badge_completion_rate = min(
                        (total_badges_earned / badge_required_to_earn) * 100, 100
                    )
                else:
                    badge_completion_rate = 100

                total_participants = participants.count()

                # 🔹 Step 3: create rows
                for idx, venue in enumerate(participants):
                    scans_per_venue = Challenge_Achiever.objects.filter(
                        challenge__venue=venue.venues,
                        scanned_at__gte=event.event_start_date,
                        scanned_at__lte=event.event_close_date
                    )

                    venue_profile = Venue_Info.objects.filter(
                        venue=venue.venues
                    ).first()

                    unique_visitors = Challenge_Achiever.objects.filter(
                        challenge__venue=venue.venues
                    ).values('customer_taken').distinct().count() or 0

                    event_rows.append({
                        "Event Title": event.title if idx == 0 else "",
                        "Total Attendees": attendees if idx == 0 else "",
                        "Participating Venues": total_participants if idx == 0 else "",
                        "Badge Completion Rates": round(badge_completion_rate, 2) if idx == 0 else "",
                        "Venue Name": venue_profile.venue_name if venue_profile else "N/A",
                        "Total Scans": scans_per_venue.count(),
                        "Unique Visitors": unique_visitors,
                    })

            # ✅ Append data
            data.extend(event_rows)

            # ✅ Track merge ranges
            row_end = row_start + len(event_rows) - 1
            event_row_ranges.append((row_start, row_end))
            row_start = row_end + 1

        # ✅ Create DataFrame
        df = pd.DataFrame(data)
        print(df)
        output = BytesIO()
        df.to_excel(output, index=False)
        output.seek(0)

        # ✅ Load workbook
        wb = load_workbook(output)
        ws = wb.active

        # ✅ Merge columns (Event Title, Attendees, Participants, Badge Rate)
        for start, end in event_row_ranges:
            excel_start = start + 2  # +1 for header, +1 for 0-index
            excel_end = end + 2

            if excel_start < excel_end:
            
                for col in [1, 2, 3, 4]:
                    ws.merge_cells(
                        start_row=excel_start,
                        start_column=col,
                        end_row=excel_end,
                        end_column=col
                    )

        # ✅ Save final file
        final_output = BytesIO()
        wb.save(final_output)
        final_output.seek(0)

        # ✅ Response
        response = HttpResponse(
            final_output,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=event_stats_report.xlsx'

        return response